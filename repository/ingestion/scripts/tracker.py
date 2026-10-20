"""
The literature to-do: repository/agent_mapping/data_tracker.xlsx, one "Qual - <Area>" tab per
agent area (the curated literature lists), kept in step with what is on disk.

Called by run.py after every write, and runnable on its own:

    uv run python repository/ingestion/scripts/tracker.py

What it writes, and only this:
  - columns J-P of every "Qual - <Area>" tab (File, Ingestion, Ingested on, Extraction,
    QC flags, Synthesis, Ingestion note), cleared and rewritten on each run, with a navy
    header so they read as the system's columns next to the user's teal ones;
  - "Do we have it?" = Yes on a row whose PDF is in repository/<topic>/raw_pdf/ (a row
    with no PDF keeps whatever the user wrote);
  - one row at the bottom of the tab for each PDF the curated list does not have, marked in
    Comments as added by ingestion (rebuilt on each run, so they never pile up; the user
    keeps one by filling A-I and deleting the mark);
  - a "Literature summary" tab with the counts, as formulas.
Every other cell — the curated columns A-I, the Notes tab, the Quant tabs — is the user's.
The workbook is the single home of the literature and data lists (since 2026-10-19; the
.md lists were folded into it). The formatting follows its own: Calibri 11, top-aligned and
wrapped, thin grey borders, Yes/No fills.

A PDF is matched to a row by year plus author surnames, with the words of the file's
description against the row's title as the tie-break. A file that matches two rows equally
is appended with a note naming them, rather than guessed.

Facts come from the folders (raw_pdf/, raw_md/, clean_md/, synthesis/), settings.yaml and
LOG.md, never from memory. The .xlsx is not in git; Dropbox's version history keeps the
previous versions.
"""

import math
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from run import LOG, NAME_RE, OBSIDIAN, REPOSITORY, scan, topics  # noqa: E402

TRACKER = REPOSITORY / "agent_mapping" / "data_tracker.xlsx"
SUMMARY = "Literature summary"
TAB = {
    "exchange_rate": "Qual - Exchange Rate", "monetary_policy": "Qual - Monetary Policy",
    "economic_activity": "Qual - Economic Activity", "fiscal_policy": "Qual - Fiscal Policy",
    "inflation": "Qual - Inflation", "labor_market": "Qual - Labor Market", "trader": "Qual - Trader",
}
BASE_HEADERS = ["Name", "Autor/Company", "Date", "Importance (1-5)", "Do we have it?", "Type", "Cluster", "Why", "Comments"]
BASE_WIDTHS = [42, 28, 12, 10, 14, 20, 28, 55, 30]
NEW_HEADERS = ["File", "Ingestion", "Ingested on", "Extraction", "QC flags", "Synthesis", "Ingestion note"]
NEW_WIDTHS = [44, 16, 15, 22, 36, 11, 40]
FIRST_NEW = len(BASE_HEADERS) + 1          # column J
ADDED_MARK = "Added by ingestion"
LABEL = {"ingested": "Ingested", "pending": "Pending", "needs_attention": "Needs attention", "on_hold": "On hold"}
_RANK = ["needs_attention", "pending", "on_hold", "ingested"]   # the worst state of a row's files wins
_STOP = {"and", "the", "of", "in", "a", "an", "for", "on", "to", "et", "al", "l1", "l2", "l3", "vs"}

# The workbook's own conventions (measured on the hand-kept tabs).
FONT = Font(name="Calibri", size=11)
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TEAL = PatternFill("solid", fgColor="418791")     # the user's header colour
NAVY = PatternFill("solid", fgColor="1F2853")     # the system's columns
EDGE = Side(style="thin", color="D7DAE2")
BORDER = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
TOP = Alignment(vertical="top", wrap_text=True)
TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
HEAD = Alignment(horizontal="center", vertical="center", wrap_text=True)
YES_FILL, NO_FILL = "DCEEE8", "FBEAE6"
STATE_FILL = {"Ingested": "DCEEE8", "Pending": "FFF1CC", "Needs attention": "F8DAD5", "On hold": "E6E8EE"}
CENTERED = {3, 4, 5, FIRST_NEW + 1, FIRST_NEW + 2, FIRST_NEW + 5}   # Date, Importance, Do we have it, Ingestion, Ingested on, Synthesis
LINE_PT = 15.0


# ── matching ─────────────────────────────────────────────────────────────────

def _norm_tokens(s: str) -> list[str]:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    # "Ch. 3" in a row's title is the "ch3" of the file name: one row per book chapter, and the
    # chapter number is what separates Gali (2015) ch. 3 from ch. 4, whose titles share words.
    s = re.sub(r"\bch(?:apter)?\.?\s*(\d+)\b", r"ch\1", s)
    return [t for t in re.split(r"[^a-z0-9]+", s) if len(t) > 1 and t not in _STOP]


# Institutions abbreviated in file names but spelled out in the lists.
_ALIASES = {"bcb": "brasil", "imf": "monetary", "bis": "settlements", "fed": "reserve"}


def _parse_stem(stem: str) -> tuple[list[str], list[str], str] | None:
    m = re.match(r"^(.*) \((.+), (\d{4})[a-z]?\)$", stem)
    if not m:
        return None
    desc, authors, year = m.groups()
    a = _norm_tokens(authors)
    return _norm_tokens(desc.replace("_", " ")), a + [_ALIASES[t] for t in a if t in _ALIASES], year


def _same_word(a: str, b: str) -> bool:
    # "rate"/"rates", "flow"/"flows" match; "international"/"intervention" do not.
    k = min(6, len(a), len(b))
    return k >= 4 and a[:k] == b[:k] or a == b


def _overlap(desc: list[str], words: list[str]) -> int:
    """How many distinct words of the title are covered by the file's description
    (counted on the title's side, so "rate" and "rates" in one description count once)."""
    return sum(1 for w in set(words) if any(_same_word(t, w) for t in desc))


def _match(file_info, rows):
    """rows: list of (row id, author tokens, year, title tokens). Returns (row, None) or (None, why)."""
    desc, authors, year = file_info
    scored = []
    for r, a_tok, r_year, n_tok in rows:
        if r_year != year:
            continue
        a = sum(1 for t in authors if t in a_tok)
        if a:
            scored.append(((a, _overlap(desc, n_tok)), r))
    if not scored:
        # No author in common ("IMF" on the file, the staff names on the row): the same year and
        # at least two title words, from one row only.
        by_title = sorted(((_overlap(desc, n_tok), r) for r, _, r_year, n_tok in rows if r_year == year), reverse=True)
        if by_title and by_title[0][0] >= 2 and (len(by_title) == 1 or by_title[0][0] > by_title[1][0]):
            return by_title[0][1], None
        return None, ""
    scored.sort(key=lambda x: x[0], reverse=True)
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        return None, "matches several rows equally: " + ", ".join(str(r) for _, r in scored[:3])
    return scored[0][1], None


def _with_synthesis(topic: str, files: list[dict]) -> set[str]:
    """File names that have a synthesis note. The vault agent names each note after its file
    (synthesis/<name>.md, since 2026-10-19; the older notes were renamed the same day), so the
    match is exact. A note named freely ("Fleming 1962 - Domestic Financial Policies…") goes to
    the one file with its year and first author whose description it shares most words with;
    a tie goes to nobody."""
    d = OBSIDIAN / topic / "synthesis"
    if not d.is_dir():
        return set()
    names = {f["name"] for f in files}
    out = {n.stem for n in d.glob("*.md") if n.stem in names}
    parsed = [(f["name"], _parse_stem(f["name"])) for f in files]
    for note in d.glob("*.md"):
        if note.stem in names or NAME_RE.match(note.stem):
            continue                     # exact, or named after a file that has no PDF here
        words = _norm_tokens(note.stem)
        scored = sorted(((_overlap(info[0], words), name) for name, info in parsed
                         if info and info[1] and info[2] in words and info[1][0] in words), reverse=True)
        if scored and scored[0][0] >= 1 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
            out.add(scored[0][1])
    return out


def _log_rows() -> dict[tuple[str, str], list[str]]:
    out = {}
    if LOG.exists():
        for line in LOG.read_text(encoding="utf-8").splitlines():
            if re.match(r"^\| \d{4}-\d{2}-\d{2} \|", line):
                cells = [c.strip() for c in line.strip("|").split("|")]
                out[(cells[1], cells[2])] = cells          # last line per file wins
    return out


# ── formatting ───────────────────────────────────────────────────────────────

def _header(ws, col: int, text: str, fill: PatternFill, width: float) -> None:
    c = ws.cell(1, col, text)
    c.font, c.fill, c.alignment, c.border = HEADER_FONT, fill, HEAD, BORDER
    ws.column_dimensions[get_column_letter(col)].width = width


def _body(cell) -> None:
    cell.font = FONT
    cell.border = BORDER
    cell.alignment = TOP_CENTER if cell.column in CENTERED else TOP


def _lines(value, width: float) -> int:
    if value in (None, ""):
        return 1
    per_line = max(int(width * 1.15), 1)     # Calibri 11 fits ~1.15 characters per width unit
    return sum(max(math.ceil(len(part) / per_line), 1) for part in str(value).split("\n"))


def _fit_height(ws, r: int, cols: range, keep_existing: bool) -> None:
    need = max(_lines(ws.cell(r, c).value, ws.column_dimensions[get_column_letter(c)].width or 10) for c in cols)
    h = need * LINE_PT + 4
    cur = ws.row_dimensions[r].height
    ws.row_dimensions[r].height = max(h, cur) if keep_existing and cur else h


def _ensure_tab(wb, title: str):
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(BASE_HEADERS, BASE_WIDTHS), 1):
        _header(ws, i, h, TEAL, w)
    ws.row_dimensions[1].height = 27.6
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    return ws


# ── one tab ──────────────────────────────────────────────────────────────────

def _fill_tab(ws, topic: str, log: dict) -> dict:
    last_new = FIRST_NEW + len(NEW_HEADERS) - 1
    # 1. remove the rows this script appended last time; clear its columns on the others
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 9).value or "").startswith(ADDED_MARK):
            ws.delete_rows(r)
            ws.row_dimensions[r].height = None
    for r in range(2, ws.max_row + 1):
        for c in range(FIRST_NEW, last_new + 1):
            cell = ws.cell(r, c)
            cell.value, cell.fill = None, PatternFill()
    for i, (h, w) in enumerate(zip(NEW_HEADERS, NEW_WIDTHS)):
        _header(ws, FIRST_NEW + i, h, NAVY, w)
    ws.cell(1, FIRST_NEW).comment = Comment(
        "The navy columns (J-P) are written by repository/ingestion/scripts/tracker.py after every "
        "ingestion; edits in them are overwritten. The teal columns (A-I) are yours.", "ingestion")

    # 2. match each PDF on disk to a row
    rows = [(r, set(_norm_tokens(ws.cell(r, 2).value)), str(ws.cell(r, 3).value or "")[:4], _norm_tokens(ws.cell(r, 1).value))
            for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    curated_rows = {r for r, *_ in rows}
    per_row: dict[int, list[dict]] = {}
    unmatched = []
    files = scan(topic)
    with_syn = _with_synthesis(topic, files)
    for f in files:
        info = _parse_stem(f["name"])
        r, why = _match(info, rows) if info else (None, "")
        if r:
            per_row.setdefault(r, []).append(f)
        else:
            unmatched.append((f, info, why))

    # 3. append a row per PDF the curated list lacks; the user moves it into the list by filling A-I
    for f, info, why in unmatched:
        r = ws.max_row + 1
        author = re.search(r"\((.+), \d{4}", f["name"])
        title = f["name"].split(" (")[0].replace("_", " ")
        values = [title[:1].upper() + title[1:], author.group(1) if author else "", int(info[2]) if info else None,
                  None, "Yes", None, None, None,
                  f"{ADDED_MARK}: not in the curated list" + (f" ({why})" if why else "")
                  + ". Fill in columns A-I and delete this sentence to keep the row."]
        for c, v in enumerate(values, 1):
            _body(ws.cell(r, c, v))
        per_row[r] = [f]

    # 4. the status columns
    for r, group in per_row.items():
        state = min((f["state"] for f in group), key=_RANK.index)
        label = LABEL[state]
        logs = [log.get((topic, f["name"])) for f in group]
        dates = sorted({lg[0] for lg in logs if lg})
        extraction = sorted({lg[3] for lg in logs if lg})
        flags = sorted({lg[7] for lg in logs if lg and lg[7] != "—"})
        if any(f["state"] == "ingested" for f in group) and not all(logs):
            dates.append("before 2026-10-19")
            extraction.append("old pipeline (x_tolerance 3)")
        n_syn = sum(1 for f in group if f["name"] in with_syn)
        has_syn = "Yes" if n_syn == len(group) else (f"{n_syn} of {len(group)}" if n_syn else "")
        notes = sorted({f["note"] for f in group if f["note"]})
        values = ["\n".join(f["name"] for f in group), label, ", ".join(dates), ", ".join(extraction),
                  "; ".join(flags), has_syn, "; ".join(notes)]
        for i, v in enumerate(values):
            _body(ws.cell(r, FIRST_NEW + i, v or None))
        ws.cell(r, FIRST_NEW + 1).fill = PatternFill("solid", fgColor=STATE_FILL[label])
        e = ws.cell(r, 5)
        e.value = "Yes"
        if e.fill.fill_type is None or (e.fill.fgColor.rgb or "").endswith(NO_FILL):
            e.fill = PatternFill("solid", fgColor=YES_FILL)    # a user's own highlight (e.g. yellow) stays
        _fit_height(ws, r, range(1, last_new + 1), keep_existing=r in curated_rows)

    # 5. curated rows without a PDF still get the grid in the new columns
    for r in curated_rows - set(per_row):
        for c in range(FIRST_NEW, last_new + 1):
            _body(ws.cell(r, c))

    ws.auto_filter.ref = f"A1:{get_column_letter(last_new)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    return {"files": sum(len(v) for v in per_row.values()), "appended": len(unmatched)}


def _summary(wb, tabs: list[tuple[str, str]]) -> None:
    """One line per area, read left to right as the work goes: the list, the PDFs, the text,
    the synthesis. Every count is of rows of the area tab, one row per source file (a book
    has a row per chapter). Dropped rows are out of the list. The two shares divide by
    Sources, so they read as how far each area is from done."""
    if SUMMARY in wb.sheetnames:
        del wb[SUMMARY]
    ws = wb.create_sheet(SUMMARY, 0)
    groups = [("", 1, 1), ("The list", 2, 3), ("Ingestion (text into the vault)", 4, 6), ("Synthesis", 7, 8)]
    heads = ["Area", "Sources", "Acquired",
             "Ingested", "% of sources", "Needs attention",
             "Synthesized", "% of sources"]
    pct_cols = {5, 8}
    # teal = the user's columns (the list), navy = the system's (ingestion, synthesis)
    fill = lambda c: TEAL if c <= 3 else NAVY  # noqa: E731
    for i, h in enumerate(heads, 1):
        for row, text in ((1, None), (2, h)):
            cell = ws.cell(row, i, text)
            cell.font, cell.alignment, cell.border, cell.fill = HEADER_FONT, HEAD, BORDER, fill(i)
        ws.column_dimensions[get_column_letter(i)].width = 24 if i == 1 else 13
    for title, a, b in groups:
        ws.cell(1, a).value = title or None
        if b > a:
            ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    first = 3
    for r, (_, tab) in enumerate(tabs, first):
        q = f"'{tab}'!"
        share = lambda col: f'=IF(B{r}=0,"",{col}{r}/B{r})'  # noqa: E731
        values = [tab.replace("Qual - ", ""),
                  f'=COUNTA({q}A:A)-1-COUNTIF({q}E:E,"Dropped")', f'=COUNTIF({q}E:E,"Yes")',
                  f'=COUNTIF({q}K:K,"Ingested")', share("D"), f'=COUNTIF({q}K:K,"Needs attention")',
                  f'=COUNTIF({q}O:O,"Yes")', share("G")]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = FONT, BORDER
            cell.alignment = Alignment(vertical="center", horizontal="left" if c == 1 else "center")
            if c in pct_cols:
                cell.number_format = "0%"
        ws.row_dimensions[r].height = 20
    t = first + len(tabs)
    ws.cell(t, 1, "Total")
    for c in range(2, len(heads) + 1):
        col = get_column_letter(c)
        if c in pct_cols:           # the total's share, not a sum of shares
            num = get_column_letter(c - 1)
            ws.cell(t, c, f'=IF(B{t}=0,"",{num}{t}/B{t})').number_format = "0%"
        else:
            ws.cell(t, c, f"=SUM({col}{first}:{col}{t - 1})")
    for c in range(1, len(heads) + 1):
        cell = ws.cell(t, c)
        cell.font, cell.border = Font(name="Calibri", size=11, bold=True), BORDER
        cell.alignment = Alignment(vertical="center", horizontal="left" if c == 1 else "center")
        cell.fill = PatternFill("solid", fgColor="EEF0F5")
    ws.row_dimensions[t].height = 20
    notes = [
        "Sources: rows of the area tab, one per file (a book has a row per chapter), dropped rows excluded. "
        "Acquired: the PDF is in the folder. Ingested: its text is in the vault. Needs attention: ingestion was tried "
        "and the tools could not read it (a scan, a broken font — the reason is in the area tab). Synthesized: the "
        "ingested file has a synthesis note. Both % columns divide by Sources.",
        "Acquired but neither Ingested nor Needs attention: not run yet (/ingest), or the area is on hold in "
        "settings.yaml (Trader). Recurring BCB documents (Copom statement, minutes, RPM) reach the vault through "
        "the Copom ETL and have no row here.",
        "The reasoning behind each area's list — scope, cluster notes, open gaps, acquisition order — is in the Notes tab.",
        "Written by repository/ingestion/scripts/tracker.py on every ingestion. In the area tabs, teal columns are yours, navy ones the system's.",
    ]
    for i, n in enumerate(notes, t + 2):
        ws.cell(i, 1, n).font = Font(name="Calibri", size=10, italic=True, color="5B6377")
    ws.freeze_panes = "B3"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100


def build() -> str:
    if not TRACKER.exists():
        raise FileNotFoundError(TRACKER)
    wb = openpyxl.load_workbook(TRACKER)
    log = _log_rows()
    done = []
    for topic in topics():
        has_pdfs = any((REPOSITORY / topic / "raw_pdf").glob("*.pdf"))
        if TAB.get(topic) not in wb.sheetnames and not has_pdfs:
            continue
        ws = _ensure_tab(wb, TAB.get(topic, f"Qual - {topic.replace('_', ' ').title()}"))
        _fill_tab(ws, topic, log)
        done.append((topic, ws.title))
    _summary(wb, done)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    tmp = TRACKER.with_name("data_tracker.tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, TRACKER)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise PermissionError(f"{TRACKER.name} is open in Excel — close it and run tracker.py again")
    return f"{TRACKER.relative_to(REPOSITORY.parent)} updated ({len(done)} areas)"


if __name__ == "__main__":
    print(build())
